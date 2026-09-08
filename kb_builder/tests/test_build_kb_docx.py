import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


SKILL_DIR = Path(__file__).resolve().parents[1]
BUILDER = SKILL_DIR / "scripts" / "build_kb_docx.py"


def document_text(doc):
    parts = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


class BuildKnowledgeBaseDocxTests(unittest.TestCase):
    def run_builder(self, payload):
        temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(temp_dir.cleanup)
        root = Path(temp_dir.name)
        source = root / "input.json"
        output = root / "output.docx"
        source.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        result = subprocess.run(
            [sys.executable, str(BUILDER), str(source), str(output)],
            capture_output=True,
            text=True,
        )
        return result, output

    def test_renders_generic_evidence_topics_and_audience_answers(self):
        payload = {
            "title": "通用知识库示例",
            "summary_columns": ["标准名", "实体类型", "一句话定义"],
            "summary_table": [
                {"标准名": "示例实体", "实体类型": "概念", "一句话定义": "用于测试结构。"}
            ],
            "sections": [
                {
                    "title": "主题一｜核心概念",
                    "entries": [
                        {
                            "标准名": "示例实体",
                            "实体类型": "概念",
                            "一句话定义": "用于测试结构。",
                            "evidence": [
                                {
                                    "claim": "示例实体用于说明通用结构。",
                                    "source": "[S] 示例规范",
                                    "location": "第 2 节",
                                    "scope": "仅作结构测试",
                                }
                            ],
                        }
                    ],
                }
            ],
            "topic_sections": [
                {
                    "title": "方案比较",
                    "intro": "项目可用同一模块表达比较、路径或组合方案。",
                    "columns": ["方案", "定位", "证据"],
                    "rows": [{"方案": "方案 A", "定位": "示例", "证据": "[G] 示例指南"}],
                }
            ],
            "qa": [
                {
                    "group": "常见问题",
                    "q": "这个结构如何使用？",
                    "answers": [
                        {"audience": "专业用户", "text": "按结构化字段审核。"},
                        {"audience": "普通用户", "text": "按主题阅读。"},
                    ],
                    "boundary": "不得超出来源。",
                    "evidence": [
                        {
                            "claim": "答案应与来源一致。",
                            "source": "[S] 示例规范",
                            "location": "第 3 节",
                            "scope": "全部受众",
                        }
                    ],
                }
            ],
            "sources": ["[S] 示例规范。"],
        }

        result, output = self.run_builder(payload)

        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue(output.exists())
        doc = Document(output)
        text = document_text(doc)
        headings = {
            paragraph.text: paragraph.style.name
            for paragraph in doc.paragraphs
            if paragraph.text
        }
        title_fonts = doc.paragraphs[0].runs[0]._element.rPr.rFonts
        self.assertEqual(title_fonts.get(qn("w:ascii")), "微软雅黑")
        self.assertEqual(title_fonts.get(qn("w:hAnsi")), "微软雅黑")
        self.assertEqual(title_fonts.get(qn("w:eastAsia")), "微软雅黑")
        self.assertEqual(headings["一、术语与实体总表"], "Heading 1")
        self.assertEqual(headings["主题一｜核心概念"], "Heading 1")
        self.assertEqual(headings["示例实体"], "Heading 2")
        self.assertEqual(headings["方案比较"], "Heading 1")
        self.assertEqual(headings["常见问题"], "Heading 2")
        self.assertEqual(headings["答案依据"], "Heading 3")
        self.assertIn("示例实体用于说明通用结构。", text)
        self.assertIn("第 2 节", text)
        self.assertIn("方案 A", text)
        self.assertIn("专业用户：按结构化字段审核。", text)
        self.assertIn("普通用户：按主题阅读。", text)
        self.assertIn("适用边界：不得超出来源。", text)
        self.assertIn("答案应与来源一致。", text)

    def test_rejects_structured_evidence_without_source(self):
        payload = {
            "title": "无来源证据",
            "sections": [
                {
                    "title": "主题",
                    "entries": [
                        {
                            "标准名": "示例实体",
                            "evidence": [{"claim": "缺少来源的结论"}],
                        }
                    ],
                }
            ],
        }

        result, output = self.run_builder(payload)

        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(output.exists())
        self.assertIn("evidence", result.stderr)

    def test_defaults_to_minimal_document_without_front_matter_or_audit_tables(self):
        payload = {
            "title": "简洁知识库",
            "subtitle": "不应默认显示的副标题",
            "meta": "不应默认显示的版本信息",
            "note": "不应默认显示的编写说明",
            "coverage_map": [
                {"一级主题": "内部覆盖审计", "二级内容定义（模版）": "内部使用", "对应知识库实体": "实体"}
            ],
            "keyword_map": [
                {"关键词维度": "内部关键词审计", "代表关键词": "关键词", "对应知识库实体": "实体"}
            ],
            "summary_columns": ["标准名", "实体类型"],
            "summary_table": [{"标准名": "实体", "实体类型": "概念"}],
        }

        result, output = self.run_builder(payload)

        self.assertEqual(result.returncode, 0, result.stderr)
        doc = Document(output)
        text = document_text(doc)
        self.assertEqual(doc.paragraphs[0].text, "简洁知识库")
        self.assertNotIn("不应默认显示的副标题", text)
        self.assertNotIn("不应默认显示的版本信息", text)
        self.assertNotIn("不应默认显示的编写说明", text)
        self.assertNotIn("内部覆盖审计", text)
        self.assertNotIn("内部关键词审计", text)
        for section in doc.sections:
            self.assertFalse(section.header.paragraphs[0].text)
            self.assertFalse(section.footer.paragraphs[0].text)


    def test_renders_antipattern_cards_and_confusable_pairs(self):
        payload = {
            "title": "反例卡与易混淆对",
            "antipatterns": [
                {
                    "id": "AP-01",
                    "缺陷代码": "X1",
                    "缺陷名称": "示例错误答法",
                    "命中模型": ["模型甲", "模型乙"],
                    "命中次数": 7,
                    "universal": True,
                    "错误表述": "这是模型说错的原话。",
                    "为什么错": "与来源不符。",
                    "正确说法": "这是有据可依的正确表述。",
                    "关联实体": ["示例实体"],
                    "来源": "[S] 示例规范",
                }
            ],
            "appendix": {
                "易混淆对": [
                    {
                        "A": "实体甲",
                        "B": "实体乙",
                        "区别维度": "适用范围",
                        "混淆后果": "会把甲的规则套到乙上。",
                        "来源": "[S] 示例规范",
                    }
                ]
            },
        }

        result, output = self.run_builder(payload)

        self.assertEqual(result.returncode, 0, result.stderr)
        text = document_text(Document(output))
        self.assertIn("常见错误答法与正确说法", text)
        self.assertIn("这是模型说错的原话。", text)
        self.assertIn("与来源不符。", text)
        self.assertIn("这是有据可依的正确表述。", text)
        self.assertIn("命中 7 次", text)
        self.assertIn("全员命中", text)
        self.assertIn("易混淆对（禁止合并）", text)
        self.assertIn("会把甲的规则套到乙上。", text)

    def test_rejects_antipattern_missing_required_field(self):
        payload = {
            "title": "缺字段的反例卡",
            "antipatterns": [
                {
                    "id": "AP-BAD",
                    "缺陷代码": "X1",
                    "错误表述": "说错的话。",
                    "为什么错": "与来源不符。",
                    "来源": "[S] 示例规范",
                }
            ],
        }

        result, _ = self.run_builder(payload)

        self.assertNotEqual(result.returncode, 0)
        self.assertIn("正确说法", result.stderr)

    def test_defect_remediation_fields_render_in_schema_order(self):
        payload = {
            "title": "缺陷弥补字段",
            "sections": [
                {
                    "title": "词条",
                    "entries": [
                        {
                            "标准名": "示例实体",
                            "实体类型": "过程",
                            "别名/同义词": "别名甲",
                            "易混淆辨析": "不要与相似实体混为一谈。",
                            "一句话定义": "用于测试字段顺序。",
                            "关键数据/证据": "示例数据。",
                            "边界与禁止表述": "不得超出来源表述。",
                            "关键安全性": "示例安全性。",
                            "风险信号与行动": "出现信号 → 采取行动。",
                            "关联实体与关系": "另一实体（适用于）",
                            "可执行下一步": "先核对状态，再提交申请。",
                            "常见问题": "示例问题。",
                            "来源": "[S] 示例规范",
                        }
                    ],
                }
            ],
        }

        result, output = self.run_builder(payload)

        self.assertEqual(result.returncode, 0, result.stderr)
        text = document_text(Document(output))
        for field in ("易混淆辨析", "边界与禁止表述", "风险信号与行动", "可执行下一步"):
            self.assertIn(field, text, f"{field} should render")

        self.assertLess(text.index("别名/同义词"), text.index("易混淆辨析"))
        self.assertLess(text.index("易混淆辨析"), text.index("一句话定义"))
        self.assertLess(text.index("关键数据/证据"), text.index("边界与禁止表述"))
        self.assertLess(text.index("关键安全性"), text.index("风险信号与行动"))
        self.assertLess(text.index("风险信号与行动"), text.index("关联实体与关系"))
        self.assertLess(text.index("关联实体与关系"), text.index("可执行下一步"))
        self.assertLess(text.index("可执行下一步"), text.index("常见问题"))


if __name__ == "__main__":
    unittest.main()
