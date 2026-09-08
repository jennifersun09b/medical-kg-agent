#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_kb_docx.py — Data-driven structured knowledge-base Word generator.

Reads ONE knowledge-base JSON file and renders a styled .docx. By default the
document starts with a single title and then goes straight into the knowledge
content: summary table, entries, topic sections, Q&A, appendix, and sources.

Domain-agnostic: works for any project. All content comes from the JSON;
this script only styles and lays it out.

Usage:
    python3 build_kb_docx.py input.json [output.docx]

If output.docx is omitted, uses the JSON's "output" field, else input basename + .docx.

Dependency: python-docx  (pip3 install python-docx)

JSON schema (all fields optional except title + at least one of summary_table/sections):
{
  "title":    "知识库标题",
  "show_front_matter": false,                  # default: title only
  "subtitle": "副标题",
  "meta":     "版本/来源/编写 一行小字",
  "note":     "开头说明(可多行, \\n 换行)",
  "cn_font":  "微软雅黑",                     # eastAsia font, default 微软雅黑
  "show_audit_sections": false,                # default: hide internal audit tables
  "summary_columns": ["标准名","实体类型","别名/同义词","一句话定义","关联实体（关系）","来源"],
  "summary_table": [ {col: value, ...}, ... ],  # each row = dict keyed by summary_columns
  "sections": [
     { "title": "二、实体词条（主题5 药物）",
       "entries": [ { "标准名": "...", "实体类型": "...", ...ordered fields... }, ... ] }
  ],
  "topic_sections": [
    {"title":"方案比较", "intro":"...", "columns":["方案","定位"], "rows":[{...}]}
  ],
  "qa": [
    {"q":"...", "a":"...", "entities":"..."},
    {"q":"...", "answers":[{"audience":"专业用户","text":"..."}],
     "boundary":"...", "evidence":[{"claim":"...","source":"..."}]}
  ],
  "appendix": {
     "实体类型": ["概念", "主体", ...],
     "关系动词": ["属于", "作用于", ...],
     "归一化规则": [ {"出现的写法":"别名","归一到":"标准名","关系类型":"同义(=)"}, ... ]
  },
  "sources": ["【S】...","【G】...", ...],
  "output": "文件名.docx"                       # optional
}

The entry card renders EVERY non-empty field in this preferred order, then any
extra fields in JSON order. Missing/empty fields are skipped.
"""
import sys, json, os

PREFERRED_FIELD_ORDER = [
    "实体类型", "别名/同义词", "易混淆辨析", "一句话定义",
    # mechanism fields — an entity's OWN mechanism; keep the drug↔entity action
    # OUT of these (that is a relationship, put it in 关联实体与关系):
    "作用机制/关键内容", "病理机制", "发生机制",
    "关键数据/证据", "边界与禁止表述",
    "适应症/适用人群", "用法用量要点", "关键安全性", "风险信号与行动",
    "关联实体与关系",
    "可执行下一步", "常见问题", "来源", "合规备注",
]

# Anti-pattern cards are worthless without the correct replacement and its
# source — a card that only says "don't say X" leaves the same gap it found.
ANTIPATTERN_REQUIRED_FIELDS = ["缺陷代码", "错误表述", "为什么错", "正确说法", "来源"]

def die(msg):
    sys.stderr.write("ERROR: " + msg + "\n"); sys.exit(1)

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    die("python-docx not installed. Run: pip3 install python-docx")

NAVY = (0x1F, 0x48, 0x77)
BLUE = (0x2E, 0x5A, 0x88)
GREY = (0x88, 0x88, 0x88)
DARK = (0x44, 0x44, 0x44)
REDISH = (0x99, 0x33, 0x33)

EVIDENCE_HEADERS = ["结论", "来源", "原文位置", "适用边界"]


def validate_evidence(items, context):
    if not isinstance(items, list):
        die("%s evidence must be a list" % context)
    for index, item in enumerate(items, 1):
        if not isinstance(item, dict):
            die("%s evidence[%d] must be an object" % (context, index))
        if not str(item.get("claim", "")).strip():
            die("%s evidence[%d] missing claim" % (context, index))
        if not str(item.get("source", "")).strip():
            die("%s evidence[%d] missing source" % (context, index))


def validate(data):
    if not str(data.get("title", "")).strip():
        die("title is required")

    for section_index, section in enumerate(data.get("sections", []), 1):
        for entry_index, entry in enumerate(section.get("entries", []), 1):
            if entry.get("evidence") is not None:
                validate_evidence(
                    entry["evidence"],
                    "sections[%d].entries[%d]" % (section_index, entry_index),
                )

    for topic_index, topic in enumerate(data.get("topic_sections", []), 1):
        if not str(topic.get("title", "")).strip():
            die("topic_sections[%d] missing title" % topic_index)
        if not isinstance(topic.get("columns"), list) or not topic["columns"]:
            die("topic_sections[%d] missing columns" % topic_index)
        if not isinstance(topic.get("rows"), list):
            die("topic_sections[%d] rows must be a list" % topic_index)

    for card_index, card in enumerate(data.get("antipatterns", []) or [], 1):
        if not isinstance(card, dict):
            die("antipatterns[%d] is not an object" % card_index)
        label = card.get("id") or "antipatterns[%d]" % card_index
        missing = [f for f in ANTIPATTERN_REQUIRED_FIELDS
                   if not str(card.get(f, "")).strip()]
        if missing:
            die("%s 缺少必填字段: %s" % (label, ", ".join(missing)))
        if card.get("evidence") is not None:
            validate_evidence(card["evidence"], label)

    for qa_index, item in enumerate(data.get("qa", []), 1):
        if item.get("evidence") is not None:
            validate_evidence(item["evidence"], "qa[%d]" % qa_index)
        if item.get("answers") is not None:
            if not isinstance(item["answers"], list) or not item["answers"]:
                die("qa[%d] answers must be a non-empty list" % qa_index)
            for answer_index, answer in enumerate(item["answers"], 1):
                if not isinstance(answer, dict) or not str(answer.get("text", "")).strip():
                    die("qa[%d] answers[%d] missing text" % (qa_index, answer_index))

def main():
    if len(sys.argv) < 2:
        die("usage: python3 build_kb_docx.py input.json [output.docx]")
    inp = sys.argv[1]
    if not os.path.exists(inp):
        die("input not found: " + inp)
    with open(inp, encoding="utf-8") as f:
        data = json.load(f)
    validate(data)

    out = sys.argv[2] if len(sys.argv) > 2 else data.get("output")
    if not out:
        out = os.path.splitext(inp)[0] + ".docx"

    CN = data.get("cn_font", "微软雅黑")
    doc = Document()

    def cn(run, size=None, bold=None, color=None):
        run.font.name = CN
        fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        for key in ("ascii", "hAnsi", "eastAsia"):
            fonts.set(qn("w:" + key), CN)
        if size: run.font.size = Pt(size)
        if bold is not None: run.font.bold = bold
        if color: run.font.color.rgb = RGBColor(*color)

    def para(text="", size=10.5, bold=False, color=None, align=None, sa=5, si=0):
        p = doc.add_paragraph()
        if align: p.alignment = align
        p.paragraph_format.space_after = Pt(sa); p.paragraph_format.space_before = Pt(si)
        cn(p.add_run(text), size, bold, color); return p

    def h(text, level=1):
        sizes = {0: 22, 1: 15, 2: 12.5, 3: 11}
        cols = {0: NAVY, 1: NAVY, 2: BLUE, 3: DARK}
        p = doc.add_heading(level=level)
        p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(5)
        cn(p.add_run(text), sizes.get(level, 12), True, cols.get(level, DARK)); return p

    def shade(cell, color):
        tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
        s.set(qn("w:fill"), color); tcPr.append(s)

    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, htext in enumerate(headers):
            c = t.rows[0].cells[i]; c.text = ""
            cn(c.paragraphs[0].add_run(str(htext)), 9, True, (0xFF, 0xFF, 0xFF)); shade(c, "2E5A88")
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = ""; cn(cells[i].paragraphs[0].add_run("" if v is None else str(v)), 8.8)
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs: p.paragraph_format.space_after = Pt(1)
        return t

    def field(name, value):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Pt(6)
        cn(p.add_run(str(name) + "："), 10, True, BLUE); cn(p.add_run(str(value)), 10)

    def evidence(items, title="逐条证据"):
        h(title, 3)
        rows = [
            [
                item.get("claim", ""),
                item.get("source", ""),
                item.get("location", ""),
                item.get("scope", ""),
            ]
            for item in items
        ]
        table(EVIDENCE_HEADERS, rows)

    def entry(d):
        title = d.get("标准名") or d.get("name") or "(未命名实体)"
        h(title, 2)
        seen = set(["标准名", "name", "evidence"])
        for k in PREFERRED_FIELD_ORDER:
            if d.get(k): field(k, d[k]); seen.add(k)
        for k, v in d.items():          # any extra fields, JSON order
            if k not in seen and v:
                field(k, v); seen.add(k)
        if d.get("evidence"):
            evidence(d["evidence"])

    # -------- opening title --------
    para(data.get("title", "知识库"), 22, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, sa=3)
    if data.get("show_front_matter") is True:
        if data.get("subtitle"):
            para(data["subtitle"], 12, False, (0x55, 0x55, 0x55), WD_ALIGN_PARAGRAPH.CENTER, sa=12)
        if data.get("meta"):
            para(data["meta"], 9, False, GREY, WD_ALIGN_PARAGRAPH.CENTER, sa=8)
        if data.get("note"):
            for line in str(data["note"]).split("\n"):
                para(line, 9.5, False, (0x66, 0x66, 0x66), sa=3)

    # Internal audit tables are omitted from the reader-facing document unless requested.
    if data.get("show_audit_sections") is True and data.get("coverage_map"):
        h(data.get("coverage_title", "本知识库对建设模版的覆盖对照"), 1)
        cmc = data.get("coverage_columns", ["一级主题", "二级内容定义（模版）", "对应知识库实体"])
        table(cmc, [[r.get(c, "") for c in cmc] for r in data["coverage_map"]])

    # -------- scenario map (use-cases → core entities/relations; 场景驱动) --------
    if data.get("scenario_map"):
        h(data.get("scenario_title", "使用场景 → 核心实体/关系"), 1)
        smc = data.get("scenario_columns", ["使用场景", "触发/问题", "核心实体链", "关键关系"])
        table(smc, [[r.get(c, "") for c in smc] for r in data["scenario_map"]])

    if data.get("show_audit_sections") is True and data.get("keyword_map"):
        h("客户关键词维度覆盖", 1)
        kmc = ["关键词维度", "代表关键词", "对应知识库实体"]
        table(kmc, [[r.get(c, "") for c in kmc] for r in data["keyword_map"]])

    # -------- summary table --------
    if data.get("summary_table"):
        cols = data.get("summary_columns",
                        ["标准名", "实体类型", "别名/同义词", "一句话定义", "关联实体（关系）", "来源"])
        h("一、术语与实体总表", 1)
        rows = [[row.get(c, "") for c in cols] for row in data["summary_table"]]
        table(cols, rows)

    # -------- sections of entries --------
    for sec in data.get("sections", []):
        if sec.get("title"): h(sec["title"], 1)
        for e in sec.get("entries", []): entry(e)

    # -------- generic topic tables --------
    # Projects can use these for comparisons, pathways, combinations, or any other topic.
    for topic in data.get("topic_sections", []):
        h(topic["title"], 1)
        if topic.get("intro"):
            para(str(topic["intro"]), 9.5, False, GREY, sa=6)
        cols = topic["columns"]
        table(cols, [[row.get(c, "") for c in cols] for row in topic["rows"]])

    # -------- anti-pattern cards --------
    # What was actually said wrong, why, and what to say instead.
    antipatterns = data.get("antipatterns") or []
    if antipatterns:
        h(data.get("antipatterns_title", "常见错误答法与正确说法"), 1)
        for card in antipatterns:
            heading = (card.get("缺陷名称") or card.get("缺陷代码")
                       or card.get("id") or "错误答法")
            code = card.get("缺陷代码", "")
            h(heading + ("（%s）" % code if code and code != heading else ""), 2)

            provenance = []
            if card.get("命中次数"):
                provenance.append("命中 %s 次" % card["命中次数"])
            if card.get("命中模型"):
                provenance.append("模型：" + "、".join(str(m) for m in card["命中模型"]))
            if card.get("universal"):
                provenance.append("全员命中")
            if provenance:
                para(" · ".join(provenance), 9, False, GREY, sa=4)

            para("❌ 错误表述：" + str(card.get("错误表述", "")), 10, False, REDISH, sa=2)
            para("为什么错：" + str(card.get("为什么错", "")), 10, sa=2)
            para("✅ 正确说法：" + str(card.get("正确说法", "")), 10, False, NAVY, sa=2)

            if card.get("关联实体"):
                entities = card["关联实体"]
                if isinstance(entities, list):
                    entities = "、".join(str(e) for e in entities)
                para("关联实体：" + str(entities), 9.5, False, DARK, sa=2)

            para("来源：" + str(card.get("来源", "")), 9.5, False, GREY, sa=6)

            if card.get("evidence"):
                evidence(card["evidence"], "证据")

    # -------- Q&A --------
    if data.get("qa"):
        h("常见问答（每问自成一段）", 1)
        if data.get("qa_intro"):
            para(data["qa_intro"], 9, False, GREY, sa=6)
        cur_group = None
        for item in data["qa"]:
            g = item.get("group")
            if g and g != cur_group:
                h(g, 2); cur_group = g
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
            cn(p.add_run("Q：" + item.get("q", "")), 10.5, True, NAVY)
            if item.get("answers"):
                for answer in item["answers"]:
                    p2 = doc.add_paragraph(); p2.paragraph_format.space_after = Pt(2)
                    audience = str(answer.get("audience", "答案")).strip() or "答案"
                    cn(p2.add_run(audience + "：" + str(answer["text"])), 10)
            else:
                p2 = doc.add_paragraph(); p2.paragraph_format.space_after = Pt(2)
                cn(p2.add_run("A：" + item.get("a", "")), 10)
            if item.get("boundary"):
                p_boundary = doc.add_paragraph(); p_boundary.paragraph_format.space_after = Pt(2)
                cn(p_boundary.add_run("适用边界：" + str(item["boundary"])), 9.5, False, REDISH)
            if item.get("entities"):
                p3 = doc.add_paragraph(); p3.paragraph_format.space_after = Pt(9)
                cn(p3.add_run("涉及实体：" + item["entities"]), 9, False, GREY)
            if item.get("evidence"):
                evidence(item["evidence"], "答案依据")

    # -------- appendix --------
    ap = data.get("appendix") or {}
    if ap:
        h("附录", 1)
        if ap.get("实体类型"):
            h("实体类型定义", 2)
            for x in ap["实体类型"]: para("· " + str(x), 10, sa=2)
        if ap.get("关系动词"):
            h("关系动词表（写关系只用这些标准动词）", 2)
            para("、".join(str(x) for x in ap["关系动词"]), 10, sa=4)
        if ap.get("关系模板"):
            h("关系模板表（按实体类型批量建关系）", 2)
            table(["主体类型", "关系", "客体类型", "示例"],
                  [[r.get("主体类型", ""), r.get("关系", ""), r.get("客体类型", ""), r.get("示例", "")] for r in ap["关系模板"]])
        if ap.get("归一化规则"):
            h("归一化规则表", 2)
            table(["出现的写法", "归一到", "关系类型"],
                  [[r.get("出现的写法", ""), r.get("归一到", ""), r.get("关系类型", "")] for r in ap["归一化规则"]])
        if ap.get("易混淆对"):
            h("易混淆对（禁止合并）", 2)
            para("以下各对实体外观相近但不等价，任何情况下不得归一为同义。",
                 9.5, False, REDISH, sa=4)
            table(["A", "B", "区别维度", "混淆后果", "来源"],
                  [[r.get("A", ""), r.get("B", ""), r.get("区别维度", ""),
                    r.get("混淆后果", ""), r.get("来源", "")] for r in ap["易混淆对"]])

    # -------- sources --------
    if data.get("sources"):
        h("来源与合规说明", 1)
        for s in data["sources"]:
            para(str(s), 9.5, False, REDISH if "合规" in str(s) or "禁区" in str(s) else None, sa=3)

    doc.save(out)
    n_entries = sum(len(s.get("entries", [])) for s in data.get("sections", []))
    stats = "  实体总表 %d 行 | 词条 %d 条 | 专题 %d 个 | 问答 %d 条" % (
        len(data.get("summary_table", [])), n_entries,
        len(data.get("topic_sections", [])), len(data.get("qa", [])))
    n_antipatterns = len(data.get("antipatterns", []) or [])
    if n_antipatterns:
        stats += " | 反例卡 %d 张" % n_antipatterns
    print("SAVED: %s" % out)
    print(stats)

if __name__ == "__main__":
    main()
